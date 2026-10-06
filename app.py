"""
FaceSync — Aplikasi Sistem Absensi berbasis Face Recognition.
File utama Flask server (app.py).

Revisi v2:
- Jadwal kuliah mingguan otomatis (menggantikan sesi manual)
- Auto-detect sesi aktif berdasarkan hari & jam
- Recognize endpoint otomatis mendeteksi sesi

Arsitektur sesuai PRD Bab 5:
- Browser menangkap gambar via webcam (getUserMedia)
- Backend (Flask) memproses face detection & recognition
- Database SQLite menyimpan data mahasiswa, pertemuan, dan absensi
"""

import os
import base64
import io
import json
import pickle
from datetime import datetime, date, time

import numpy as np
from PIL import Image
from flask import Flask, request, jsonify, send_file, send_from_directory
from flask_cors import CORS
from sqlalchemy import inspect as sa_inspect, text

from config import Config
from models import db, Jadwal, Mahasiswa, Pertemuan, Absensi

# Coba import face_recognition (membutuhkan dlib terinstall)
try:
    import face_recognition
    FACE_RECOGNITION_AVAILABLE = True
except ImportError:
    FACE_RECOGNITION_AVAILABLE = False
    print("[WARNING] Library face_recognition tidak ditemukan.")
    print("         Jalankan: pip install face_recognition")
    print("         Sistem akan berjalan tanpa fitur pengenalan wajah.")


# ============================================================
# Inisialisasi Aplikasi Flask
# ============================================================

app = Flask(__name__, static_folder='.', static_url_path='')
app.config.from_object(Config)
app.config['SEND_FILE_MAX_AGE_DEFAULT'] = 0  # Jangan cache file statis
CORS(app)

@app.after_request
def add_no_cache_headers(response):
    """Paksa browser untuk selalu mengunduh ulang file CSS/JS terbaru."""
    if 'text/css' in response.content_type or 'javascript' in response.content_type or 'text/html' in response.content_type:
        response.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
        response.headers['Pragma'] = 'no-cache'
        response.headers['Expires'] = '0'
    return response

db.init_app(app)

# Buat folder upload & export jika belum ada
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
os.makedirs(app.config['EXPORT_FOLDER'], exist_ok=True)


# Buat tabel database & jalankan migrasi
with app.app_context():
    db.create_all()
    # Migrasi: tambahkan kolom jadwal_id ke tabel pertemuan jika belum ada
    try:
        inspector = sa_inspect(db.engine)
        columns = [c['name'] for c in inspector.get_columns('pertemuan')]
        if 'jadwal_id' not in columns:
            with db.engine.connect() as conn:
                conn.execute(text("ALTER TABLE pertemuan ADD COLUMN jadwal_id INTEGER REFERENCES jadwal(id)"))
                conn.commit()
            print("[MIGRASI] Kolom jadwal_id berhasil ditambahkan ke tabel pertemuan.")
    except Exception as e:
        print(f"[MIGRASI] Info: {e}")


# ============================================================
# Helper Functions
# ============================================================

def decode_base64_image(base64_string):
    """Mengubah string base64 menjadi numpy array (image)."""
    # Bersihkan prefix data URI jika ada
    if ',' in base64_string:
        base64_string = base64_string.split(',')[1]

    image_bytes = base64.b64decode(base64_string)
    image = Image.open(io.BytesIO(image_bytes)).convert('RGB')
    return np.array(image)


def save_base64_image(base64_string, filepath):
    """Menyimpan gambar base64 ke file."""
    if ',' in base64_string:
        base64_string = base64_string.split(',')[1]

    image_bytes = base64.b64decode(base64_string)
    with open(filepath, 'wb') as f:
        f.write(image_bytes)


def ensure_today_sessions():
    """
    Buat pertemuan untuk SEMUA jadwal hari ini sekaligus.
    Dipanggil saat halaman riwayat atau active-session diakses,
    agar semua matkul hari ini muncul meskipun belum masuk jam-nya.
    """
    now = datetime.now()
    today = now.date()
    current_day = now.weekday()  # 0=Senin, 1=Selasa, ..., 6=Minggu

    jadwal_list = Jadwal.query.filter_by(hari=current_day).all()
    created = 0
    for j in jadwal_list:
        existing = Pertemuan.query.filter_by(jadwal_id=j.id, tanggal=today).first()
        if not existing:
            pertemuan = Pertemuan(
                judul=j.mata_kuliah,
                tanggal=today,
                jam_mulai=j.jam_mulai,
                jam_selesai=j.jam_selesai,
                status_sesi='buka',
                jadwal_id=j.id
            )
            db.session.add(pertemuan)
            created += 1
    if created > 0:
        db.session.commit()
    return created


def get_active_jadwal_and_session():
    """
    Cek hari & waktu saat ini, temukan jadwal yang cocok,
    dan otomatis buat/buka pertemuan jika belum ada.
    Returns: (jadwal, pertemuan) atau (None, None)
    """
    # Pastikan semua pertemuan hari ini sudah dibuat
    ensure_today_sessions()

    now = datetime.now()
    current_day = now.weekday()
    current_time = now.time()

    # Cari jadwal yang cocok dengan hari & waktu saat ini
    jadwal_list = Jadwal.query.filter_by(hari=current_day).all()
    active_jadwal = None
    for j in jadwal_list:
        if j.jam_mulai <= current_time <= j.jam_selesai:
            active_jadwal = j
            break

    if not active_jadwal:
        return None, None

    # Ambil pertemuan yang sudah pasti ada (dibuat oleh ensure_today_sessions)
    today = now.date()
    pertemuan = Pertemuan.query.filter_by(jadwal_id=active_jadwal.id, tanggal=today).first()

    if pertemuan and pertemuan.status_sesi == 'tutup':
        # Buka kembali jika masih dalam waktu
        pertemuan.status_sesi = 'buka'
        db.session.commit()

    return active_jadwal, pertemuan


# ============================================================
# Route: Serve Frontend
# ============================================================

@app.route('/')
def serve_index():
    """Menyajikan halaman utama (index.html)."""
    return send_from_directory('.', 'index.html')


@app.route('/admin')
def serve_admin():
    """Menyajikan halaman admin (admin.html)."""
    return send_from_directory('.', 'admin.html')


# ============================================================
# API: Jadwal Kuliah (CRUD) — Revisi v2
# ============================================================

@app.route('/api/jadwal', methods=['GET'])
def get_jadwal():
    """Mengambil seluruh jadwal kuliah, diurutkan per hari lalu jam."""
    jadwal_list = Jadwal.query.order_by(Jadwal.hari, Jadwal.jam_mulai).all()
    return jsonify({
        'success': True,
        'data': [j.to_dict() for j in jadwal_list],
        'total': len(jadwal_list)
    })


@app.route('/api/jadwal', methods=['POST'])
def create_jadwal():
    """Membuat jadwal kuliah baru."""
    data = request.get_json()

    hari = data.get('hari')
    jam_mulai_str = data.get('jam_mulai', '')
    jam_selesai_str = data.get('jam_selesai', '')
    mata_kuliah = data.get('mata_kuliah', '').strip()
    dosen = data.get('dosen', '').strip()
    ruangan = data.get('ruangan', '').strip()

    if hari is None or not jam_mulai_str or not jam_selesai_str or not mata_kuliah:
        return jsonify({'success': False, 'message': 'Hari, jam mulai, jam selesai, dan mata kuliah wajib diisi.'}), 400

    try:
        hari = int(hari)
        jam_mulai = datetime.strptime(jam_mulai_str, '%H:%M').time()
        jam_selesai = datetime.strptime(jam_selesai_str, '%H:%M').time()
    except (ValueError, TypeError):
        return jsonify({'success': False, 'message': 'Format hari atau waktu tidak valid.'}), 400

    if hari < 0 or hari > 6:
        return jsonify({'success': False, 'message': 'Hari harus antara 0 (Senin) sampai 6 (Minggu).'}), 400

    new_jadwal = Jadwal(
        hari=hari,
        jam_mulai=jam_mulai,
        jam_selesai=jam_selesai,
        mata_kuliah=mata_kuliah,
        dosen=dosen or None,
        ruangan=ruangan or None
    )
    db.session.add(new_jadwal)
    db.session.commit()

    return jsonify({
        'success': True,
        'message': f'Jadwal "{mata_kuliah}" berhasil ditambahkan.',
        'data': new_jadwal.to_dict()
    }), 201


@app.route('/api/jadwal/<int:jadwal_id>', methods=['PUT'])
def update_jadwal(jadwal_id):
    """Update jadwal kuliah."""
    jadwal = Jadwal.query.get_or_404(jadwal_id)
    data = request.get_json()

    if 'hari' in data:
        jadwal.hari = int(data['hari'])
    if 'jam_mulai' in data:
        jadwal.jam_mulai = datetime.strptime(data['jam_mulai'], '%H:%M').time()
    if 'jam_selesai' in data:
        jadwal.jam_selesai = datetime.strptime(data['jam_selesai'], '%H:%M').time()
    if 'mata_kuliah' in data:
        jadwal.mata_kuliah = data['mata_kuliah'].strip()
    if 'dosen' in data:
        jadwal.dosen = data['dosen'].strip() or None
    if 'ruangan' in data:
        jadwal.ruangan = data['ruangan'].strip() or None

    db.session.commit()
    return jsonify({'success': True, 'message': 'Jadwal berhasil diperbarui.', 'data': jadwal.to_dict()})


@app.route('/api/jadwal/<int:jadwal_id>', methods=['DELETE'])
def delete_jadwal(jadwal_id):
    """Menghapus jadwal kuliah."""
    jadwal = Jadwal.query.get_or_404(jadwal_id)
    mata_kuliah = jadwal.mata_kuliah
    db.session.delete(jadwal)
    db.session.commit()
    return jsonify({'success': True, 'message': f'Jadwal "{mata_kuliah}" berhasil dihapus.'})


# ============================================================
# API: Sesi Aktif Otomatis — Revisi v2
# ============================================================

@app.route('/api/active-session', methods=['GET'])
def get_active_session():
    """
    Cek jadwal yang sedang berlangsung saat ini.
    Otomatis membuat pertemuan baru jika belum ada untuk hari ini.
    Digunakan oleh halaman mahasiswa untuk auto-detect kelas.
    """
    jadwal, pertemuan = get_active_jadwal_and_session()

    if not jadwal or not pertemuan:
        # Tidak ada kelas saat ini, cari jadwal berikutnya
        now = datetime.now()
        current_day = now.weekday()
        current_time = now.time()

        next_jadwal = None
        # Cari di hari yang sama (setelah waktu saat ini)
        today_remaining = Jadwal.query.filter(
            Jadwal.hari == current_day,
            Jadwal.jam_mulai > current_time
        ).order_by(Jadwal.jam_mulai).first()

        if today_remaining:
            next_jadwal = today_remaining
        else:
            # Cari di hari-hari berikutnya
            for offset in range(1, 8):
                next_day = (current_day + offset) % 7
                first_on_day = Jadwal.query.filter_by(hari=next_day).order_by(Jadwal.jam_mulai).first()
                if first_on_day:
                    next_jadwal = first_on_day
                    break

        response = {
            'success': True,
            'active': False,
            'message': 'Tidak ada kelas yang berlangsung saat ini.'
        }
        if next_jadwal:
            response['next_jadwal'] = next_jadwal.to_dict()

        return jsonify(response)

    hadir_count = Absensi.query.filter_by(pertemuan_id=pertemuan.id, status='hadir').count()
    total_students = Mahasiswa.query.count()

    return jsonify({
        'success': True,
        'active': True,
        'session': pertemuan.to_dict(),
        'jadwal': jadwal.to_dict(),
        'hadir_count': hadir_count,
        'total_students': total_students
    })


# ============================================================
# API: Mahasiswa (CRUD + Registrasi Wajah) — PRD Fase 1
# ============================================================

@app.route('/api/students', methods=['GET'])
def get_students():
    """Mengambil daftar seluruh mahasiswa."""
    students = Mahasiswa.query.order_by(Mahasiswa.created_at.desc()).all()
    return jsonify({
        'success': True,
        'data': [s.to_dict() for s in students],
        'total': len(students)
    })


@app.route('/api/students/<int:student_id>', methods=['GET'])
def get_student(student_id):
    """Mengambil data satu mahasiswa berdasarkan ID."""
    student = Mahasiswa.query.get_or_404(student_id)
    return jsonify({'success': True, 'data': student.to_dict()})


@app.route('/api/students/register', methods=['POST'])
def register_student():
    """
    Registrasi mahasiswa baru + Face Encoding.
    PRD Fase 1: Mahasiswa mengisi NIM & nama, lalu mengambil 3-5 foto wajah
    via webcam untuk di-generate menjadi face encoding.
    """
    data = request.get_json()

    nim = data.get('nim', '').strip()
    nama = data.get('nama', '').strip()
    kelas = data.get('kelas', '').strip()
    images = data.get('images', [])  # Array of base64 strings

    # Validasi input
    if not nim or not nama or not kelas:
        return jsonify({'success': False, 'message': 'NIM, Nama, dan Kelas wajib diisi.'}), 400

    if len(images) < 3:
        return jsonify({'success': False, 'message': 'Minimal 3 foto wajah diperlukan.'}), 400

    # Cek duplikasi NIM
    existing = Mahasiswa.query.filter_by(nim=nim).first()
    if existing:
        return jsonify({'success': False, 'message': f'Mahasiswa dengan NIM {nim} sudah terdaftar.'}), 409

    # Proses Face Encoding
    all_encodings = []
    foto_dir = os.path.join(app.config['UPLOAD_FOLDER'], nim)
    os.makedirs(foto_dir, exist_ok=True)

    for i, img_base64 in enumerate(images):
        # Simpan foto ke disk
        filepath = os.path.join(foto_dir, f'face_{i + 1}.jpg')
        save_base64_image(img_base64, filepath)

        if FACE_RECOGNITION_AVAILABLE:
            # Decode gambar dan generate face encoding
            image_array = decode_base64_image(img_base64)
            encodings = face_recognition.face_encodings(image_array)

            if encodings:
                all_encodings.append(encodings[0])

    # Hitung rata-rata encoding untuk akurasi lebih baik
    final_encoding = None
    if FACE_RECOGNITION_AVAILABLE and all_encodings:
        final_encoding = np.mean(all_encodings, axis=0).tolist()

    # Simpan ke database
    new_student = Mahasiswa(
        nim=nim,
        nama=nama,
        kelas=kelas,
        face_encoding=final_encoding,
        foto_path=foto_dir
    )
    db.session.add(new_student)
    db.session.commit()

    encoding_count = len(all_encodings)
    total_images = len(images)

    return jsonify({
        'success': True,
        'message': f'Mahasiswa {nama} berhasil didaftarkan. ({encoding_count}/{total_images} wajah berhasil di-encode)',
        'data': new_student.to_dict()
    }), 201


@app.route('/api/students/<int:student_id>', methods=['PUT'])
def update_student(student_id):
    """Update data mahasiswa (NIM, Nama, Kelas)."""
    student = Mahasiswa.query.get_or_404(student_id)
    data = request.get_json()

    if 'nim' in data:
        # Cek duplikasi NIM jika berubah
        if data['nim'] != student.nim:
            existing = Mahasiswa.query.filter_by(nim=data['nim']).first()
            if existing:
                return jsonify({'success': False, 'message': 'NIM sudah digunakan mahasiswa lain.'}), 409
        student.nim = data['nim']
    if 'nama' in data:
        student.nama = data['nama']
    if 'kelas' in data:
        student.kelas = data['kelas']

    db.session.commit()
    return jsonify({'success': True, 'message': 'Data mahasiswa berhasil diperbarui.', 'data': student.to_dict()})


@app.route('/api/students/<int:student_id>', methods=['DELETE'])
def delete_student(student_id):
    """Menghapus data mahasiswa beserta file foto."""
    student = Mahasiswa.query.get_or_404(student_id)

    # Hapus folder foto jika ada
    if student.foto_path and os.path.exists(student.foto_path):
        import shutil
        shutil.rmtree(student.foto_path, ignore_errors=True)

    db.session.delete(student)
    db.session.commit()
    return jsonify({'success': True, 'message': f'Mahasiswa {student.nama} berhasil dihapus.'})


# ============================================================
# API: Pertemuan / Riwayat Sesi — PRD Fase 1 & 2
# ============================================================

@app.route('/api/sessions', methods=['GET'])
def get_sessions():
    """Mengambil daftar seluruh pertemuan (termasuk yang dibuat otomatis)."""
    # Pastikan semua pertemuan hari ini sudah dibuat
    ensure_today_sessions()

    sessions = Pertemuan.query.order_by(Pertemuan.tanggal.desc(), Pertemuan.jam_mulai.desc()).all()
    result = []
    for session in sessions:
        session_dict = session.to_dict()
        # Hitung jumlah yang hadir di sesi ini
        hadir_count = Absensi.query.filter_by(pertemuan_id=session.id, status='hadir').count()
        total_students = Mahasiswa.query.count()
        session_dict['hadir_count'] = hadir_count
        session_dict['total_students'] = total_students
        result.append(session_dict)

    return jsonify({'success': True, 'data': result, 'total': len(result)})


@app.route('/api/sessions/<int:session_id>', methods=['DELETE'])
def delete_session(session_id):
    """Menghapus pertemuan beserta catatan absensinya."""
    session = Pertemuan.query.get_or_404(session_id)
    db.session.delete(session)
    db.session.commit()
    return jsonify({'success': True, 'message': f'Pertemuan "{session.judul}" berhasil dihapus.'})


# ============================================================
# API: Face Recognition & Absensi — PRD Fase 2 (Core)
# ============================================================

@app.route('/api/recognize', methods=['POST'])
def recognize_face():
    """
    Endpoint utama pengenalan wajah.
    Revisi v2: Otomatis mendeteksi sesi aktif dari jadwal jika session_id tidak diberikan.

    Menerima frame base64 dari webcam browser, mendeteksi wajah,
    mencocokkan dengan face encoding tersimpan, dan mencatat kehadiran.
    """
    data = request.get_json()
    image_base64 = data.get('image', '')
    session_id = data.get('session_id')  # Opsional di v2

    if not image_base64:
        return jsonify({'success': False, 'message': 'Image diperlukan.'}), 400

    # Jika session_id tidak diberikan, auto-detect dari jadwal
    if not session_id:
        _, pertemuan = get_active_jadwal_and_session()
        if not pertemuan:
            return jsonify({'success': False, 'message': 'Tidak ada sesi aktif saat ini.'}), 404
        session_id = pertemuan.id
    else:
        # Validasi sesi yang diberikan
        pertemuan = Pertemuan.query.get(session_id)
        if not pertemuan:
            return jsonify({'success': False, 'message': 'Sesi tidak ditemukan.'}), 404
        if pertemuan.status_sesi != 'buka':
            return jsonify({'success': False, 'message': 'Sesi absensi belum dibuka.'}), 403

    if not FACE_RECOGNITION_AVAILABLE:
        return jsonify({'success': False, 'message': 'Library face_recognition tidak tersedia di server.'}), 503

    # Decode image
    try:
        image_array = decode_base64_image(image_base64)
    except Exception as e:
        return jsonify({'success': False, 'message': f'Gagal memproses gambar: {str(e)}'}), 400

    # Deteksi wajah dalam frame
    face_locations = face_recognition.face_locations(image_array)
    face_encodings_in_frame = face_recognition.face_encodings(image_array, face_locations)

    if not face_encodings_in_frame:
        return jsonify({
            'success': True,
            'recognized': False,
            'message': 'Tidak ada wajah terdeteksi dalam frame.'
        })

    # Ambil semua face encoding dari database
    students = Mahasiswa.query.filter(Mahasiswa.face_encoding.isnot(None)).all()
    if not students:
        return jsonify({
            'success': True,
            'recognized': False,
            'message': 'Belum ada data wajah mahasiswa terdaftar.'
        })

    known_encodings = []
    known_students = []
    for student in students:
        encoding = student.face_encoding
        if encoding is not None:
            known_encodings.append(np.array(encoding))
            known_students.append(student)

    results = []
    tolerance = app.config['FACE_RECOGNITION_TOLERANCE']

    for face_encoding in face_encodings_in_frame:
        # Hitung jarak ke semua wajah yang terdaftar
        distances = face_recognition.face_distance(known_encodings, face_encoding)
        best_match_idx = np.argmin(distances)
        best_distance = distances[best_match_idx]

        if best_distance <= tolerance:
            matched_student = known_students[best_match_idx]
            confidence = float(round((1 - best_distance) * 100, 2))

            # PRD: Pencegahan Duplikasi Absen
            # Cek apakah mahasiswa sudah tercatat hadir di sesi ini
            existing_record = Absensi.query.filter_by(
                mahasiswa_id=matched_student.id,
                pertemuan_id=session_id,
                status='hadir'
            ).first()

            if existing_record:
                results.append({
                    'recognized': True,
                    'already_recorded': True,
                    'student': matched_student.to_dict(),
                    'confidence': confidence,
                    'message': f'{matched_student.nama} sudah tercatat hadir.'
                })
            else:
                # Catat kehadiran baru
                new_absensi = Absensi(
                    mahasiswa_id=matched_student.id,
                    pertemuan_id=session_id,
                    waktu_absen=datetime.now(),
                    status='hadir',
                    confidence_score=confidence
                )
                db.session.add(new_absensi)
                db.session.commit()

                results.append({
                    'recognized': True,
                    'already_recorded': False,
                    'student': matched_student.to_dict(),
                    'confidence': confidence,
                    'message': f'{matched_student.nama} berhasil dicatat hadir!'
                })
        else:
            results.append({
                'recognized': False,
                'message': 'Wajah tidak dikenali.',
                'distance': round(float(best_distance), 4)
            })

    return jsonify({'success': True, 'results': results})


# ============================================================
# API: Dashboard & Statistik — PRD Fase 3
# ============================================================

@app.route('/api/dashboard/stats', methods=['GET'])
def get_dashboard_stats():
    """Statistik ringkasan untuk halaman Dashboard."""
    total_students = Mahasiswa.query.count()
    total_sessions = Pertemuan.query.count()
    total_jadwal = Jadwal.query.count()

    # Hitung rata-rata kehadiran
    if total_sessions > 0 and total_students > 0:
        total_hadir = Absensi.query.filter_by(status='hadir').count()
        avg_attendance = round((total_hadir / (total_students * total_sessions)) * 100, 1)
    else:
        avg_attendance = 0

    return jsonify({
        'success': True,
        'data': {
            'total_students': total_students,
            'total_sessions': total_sessions,
            'total_jadwal': total_jadwal,
            'avg_attendance': avg_attendance
        }
    })


@app.route('/api/dashboard/recent', methods=['GET'])
def get_recent_sessions():
    """Mengambil 5 pertemuan terakhir untuk Dashboard."""
    sessions = Pertemuan.query.order_by(Pertemuan.tanggal.desc(), Pertemuan.jam_mulai.desc()).limit(5).all()
    result = []
    total_students = Mahasiswa.query.count()

    for session in sessions:
        session_dict = session.to_dict()
        hadir_count = Absensi.query.filter_by(pertemuan_id=session.id, status='hadir').count()
        session_dict['hadir_count'] = hadir_count
        session_dict['total_students'] = total_students
        result.append(session_dict)

    return jsonify({'success': True, 'data': result})


# ============================================================
# API: Riwayat Absensi Mandiri — PRD Fase 3
# ============================================================

@app.route('/api/history/<nim>', methods=['GET'])
def get_student_history(nim):
    """
    Cek riwayat kehadiran mahasiswa berdasarkan NIM.
    PRD Fase 3: Mahasiswa memasukkan NIM untuk melihat status kehadirannya.
    """
    student = Mahasiswa.query.filter_by(nim=nim).first()
    if not student:
        return jsonify({'success': False, 'message': f'Mahasiswa dengan NIM {nim} tidak ditemukan.'}), 404

    # Ambil semua pertemuan
    all_sessions = Pertemuan.query.order_by(Pertemuan.tanggal.desc()).all()
    total_sessions = len(all_sessions)

    # Ambil catatan absensi mahasiswa ini
    records = Absensi.query.filter_by(mahasiswa_id=student.id).all()
    attended_session_ids = {r.pertemuan_id for r in records if r.status == 'hadir'}
    total_hadir = len(attended_session_ids)

    # Buat riwayat lengkap (hadir + alpa)
    history = []
    for session in all_sessions:
        absensi_record = Absensi.query.filter_by(
            mahasiswa_id=student.id,
            pertemuan_id=session.id
        ).first()

        history.append({
            'tanggal': session.tanggal.isoformat() if session.tanggal else None,
            'pertemuan_judul': session.judul,
            'waktu_absen': absensi_record.waktu_absen.strftime('%H:%M WIB') if absensi_record and absensi_record.waktu_absen else '-',
            'status': absensi_record.status if absensi_record else 'alpa',
            'metode': 'Pengenalan Wajah' if absensi_record else '-'
        })

    # Hitung persentase
    rate = round((total_hadir / total_sessions * 100), 1) if total_sessions > 0 else 0

    return jsonify({
        'success': True,
        'data': {
            'student': student.to_dict(),
            'total_sessions': total_sessions,
            'total_hadir': total_hadir,
            'rate': rate,
            'history': history
        }
    })


# ============================================================
# API: Ekspor Laporan — PRD Fase 3
# ============================================================

@app.route('/api/export/<int:session_id>', methods=['GET'])
def export_session_report(session_id):
    """
    Ekspor rekap kehadiran per pertemuan ke file Excel.
    PRD Fase 3: Dosen mengekspor laporan kehadiran.
    """
    session = Pertemuan.query.get_or_404(session_id)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        return jsonify({'success': False, 'message': 'Library openpyxl tidak tersedia.'}), 503

    wb = Workbook()
    ws = wb.active
    ws.title = 'Rekap Kehadiran'

    # Styling
    header_font = Font(bold=True, color='FFFFFF', size=12)
    header_fill = PatternFill(start_color='6366F1', end_color='6366F1', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    # Judul
    ws.merge_cells('A1:E1')
    ws['A1'] = f'Rekap Kehadiran — {session.judul}'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = f'Tanggal: {session.tanggal}'
    ws['A3'] = f'Waktu: {session.jam_mulai.strftime("%H:%M")} - {session.jam_selesai.strftime("%H:%M")}'

    # Header tabel
    headers = ['No', 'NIM', 'Nama', 'Status', 'Waktu Absen']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    # Data
    all_students = Mahasiswa.query.order_by(Mahasiswa.nim).all()
    for idx, student in enumerate(all_students, 1):
        row = idx + 5
        absensi = Absensi.query.filter_by(
            mahasiswa_id=student.id,
            pertemuan_id=session_id
        ).first()

        status = absensi.status.capitalize() if absensi else 'Alpa'
        waktu = absensi.waktu_absen.strftime('%H:%M:%S') if absensi and absensi.waktu_absen else '-'

        ws.cell(row=row, column=1, value=idx).border = thin_border
        ws.cell(row=row, column=2, value=student.nim).border = thin_border
        ws.cell(row=row, column=3, value=student.nama).border = thin_border

        status_cell = ws.cell(row=row, column=4, value=status)
        status_cell.border = thin_border
        if status == 'Hadir':
            status_cell.fill = PatternFill(start_color='D1FAE5', fill_type='solid')
        else:
            status_cell.fill = PatternFill(start_color='FEE2E2', fill_type='solid')

        ws.cell(row=row, column=5, value=waktu).border = thin_border

    # Auto-width kolom
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[col_letter].width = max_length + 4

    # Simpan file
    filename = f'rekap_{session.judul.replace(" ", "_")}_{session.tanggal}.xlsx'
    filepath = os.path.join(app.config['EXPORT_FOLDER'], filename)
    wb.save(filepath)

    return send_file(filepath, as_attachment=True, download_name=filename)


# ============================================================
# API: Absensi per Sesi (untuk detail view)
# ============================================================

@app.route('/api/sessions/<int:session_id>/attendance', methods=['GET'])
def get_session_attendance(session_id):
    """Mengambil daftar kehadiran untuk sesi tertentu."""
    session = Pertemuan.query.get_or_404(session_id)
    records = Absensi.query.filter_by(pertemuan_id=session_id).order_by(Absensi.waktu_absen.desc()).all()

    return jsonify({
        'success': True,
        'session': session.to_dict(),
        'data': [r.to_dict() for r in records],
        'total_hadir': len([r for r in records if r.status == 'hadir'])
    })


# ============================================================
# Jalankan Server
# ============================================================

if __name__ == '__main__':
    print("\n" + "=" * 50)
    print("  FaceSync - Sistem Absensi Pintar v2")
    print("  Face Recognition: ", "Aktif [OK]" if FACE_RECOGNITION_AVAILABLE else "Tidak Tersedia [X]")
    print("  Jadwal Otomatis  : Aktif [OK]")
    print("=" * 50)
    print(f"  Server berjalan di: http://localhost:5000")
    print("=" * 50 + "\n")

    app.run(debug=True, host='0.0.0.0', port=5000)
